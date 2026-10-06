from app.xlsx_export import write
from openpyxl import load_workbook
def test_xlsx_export_has_real_workbook(tmp_path):
 p=write(tmp_path/"r.xlsx",{"Observations":[{"a":1,"api_key":"x"}]});w=load_workbook(p);assert "Observations" in w.sheetnames and w["Observations"]["B2"].value=="***"
