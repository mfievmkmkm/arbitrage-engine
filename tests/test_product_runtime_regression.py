import asyncio
import json
import zipfile
import time
from types import SimpleNamespace
import aiosqlite
from openpyxl import load_workbook
from app.audit_export import build as export
from app.db import Diary
from app.strategy_diary import init as strategies_init, record
from app.ledger_store import init as ledger_init
from app.spot_future_cycle_service import CycleService
from app.spot_future_paper_engine import Engine
from app.strategy_paper_coordinator import Coordinator
from app.spot_future_paper_store import init as sf_init
from app.paper_ledger import restore
from app.paper import PaperEngine
from app.bankroll_ledger import Ledger
from app.risk import RiskGuard
from app.engine import Scanner
from app.discovery import RotatingUniverse
from app.instruments import from_market
from app.zerox_research import Provider


def sf_op():
    return dict(
        base="X",
        exchange="a",
        direction="LONG_SPOT_SHORT_FUTURE",
        notional=5,
        base_qty=0.05,
        prices=dict(spot_buy=100, spot_sell=99, future_buy=110, future_sell=109),
        fee_pct=0,
        safety_pct=0,
        funding_pct=0,
        hypothetical_edge=3,
    )


def test_secondary_positions_and_ids_survive_restart(tmp_path):
    async def go():
        path = str(tmp_path / "a.db")

        class Service:
            source = SimpleNamespace(watch_pairs=set())

            async def cycle(self):
                return [sf_op()], []

        first = CycleService(Service(), Coordinator(Engine()), 2, path)
        await first.restore()
        await first.cycle()
        second = CycleService(Service(), Coordinator(Engine()), 2, path)
        await second.restore()
        assert second.paper.sf.positions[1].base == "X" and second.paper.sf.next_id == 2
        second.entry_enabled = False
        await second.cycle()
        assert len(second.paper.sf.positions) == 1
        async with aiosqlite.connect(path) as d:
            assert (
                await (
                    await d.execute("SELECT COUNT(*) FROM spot_future_marks")
                ).fetchone()
            )[0] == 2

    asyncio.run(go())


def test_complete_audit_export_is_readable_and_redacts_payload(tmp_path):
    async def go():
        path = str(tmp_path / "a.db")
        d = Diary(path)
        await d.init()
        await strategies_init(path)
        await ledger_init(path)
        await sf_init(path)
        await record(
            path,
            [
                dict(
                    ts=1,
                    strategy="x",
                    symbol="X",
                    venue="a",
                    edge=2,
                    notional=5,
                    payload={"api_key": "do-not-export", "v": 1},
                )
            ],
        )
        await d.record_decisions(
            [
                dict(
                    ts=1,
                    strategy="x",
                    symbol="X",
                    buy="a",
                    sell="b",
                    action="SKIP",
                    reason="LIMIT",
                )
            ]
        )
        xlsx, archive = await export(path, tmp_path / "export")
        wb = load_workbook(xlsx)
        assert "Decisions" in wb.sheetnames and "Paper marks" in wb.sheetnames
        with zipfile.ZipFile(archive) as z:
            text = z.read("strategy_observations.csv").decode()
            assert "do-not-export" not in text and "***" in text
            assert "LIMIT" in z.read("decisions.csv").decode()

    asyncio.run(go())


def test_scanner_keeps_a_watched_route_when_entry_edge_disappears():
    async def go():
        symbol = "X/USDT:USDT"

        class Client:
            async def fetch_order_book(self, symbol, limit=20):
                return {"bids": [[99, 100]], "asks": [[100, 100]]}

        scanner = Scanner(["binance", "bybit"], 5, 12)
        scanner.clients = {v: Client() for v in scanner.ids}
        scanner.symbols = {v: {symbol} for v in scanner.ids}
        market = {
            "symbol": symbol,
            "base": "X",
            "quote": "USDT",
            "settle": "USDT",
            "contract": True,
            "linear": True,
            "contractSize": 1,
        }
        scanner.specs = {v: {symbol: from_market(v, market)} for v in scanner.ids}
        scanner.universe = RotatingUniverse(scanner.symbols)
        scanner.watch_routes = {(symbol, "binance", "bybit")}
        rows = await scanner.scan()
        assert len(rows) == 1 and rows[0]["hypothetical_edge"] < 0
        scanner.paused = True
        rows = await scanner.scan()
        assert len(rows) == 1 and rows[0]["hypothetical_edge"] < 0
        scanner.watch_routes = set()
        assert await scanner.scan() == []

    asyncio.run(go())


def test_projected_future_funding_is_not_paper_profit(tmp_path):
    async def go():
        d = Diary(str(tmp_path / "a.db"))
        await d.init()
        e = PaperEngine(d)
        op = dict(
            symbol="X",
            buy="a",
            sell="b",
            notional=5,
            entry_buy=100,
            entry_sell=110,
            executable=10,
            exit_buy=100,
            exit_sell=110,
            exit_spread=10,
            fee_pct=0,
            funding_pct=99,
        )
        position = await e.open(op)
        await e.mark_and_exit([op])
        assert position.current_net_usd == 0

    asyncio.run(go())


def test_daily_loss_is_restored_from_closed_positions(tmp_path):
    async def go():
        d = Diary(str(tmp_path / "a.db"))
        await d.init()
        async with aiosqlite.connect(d.path) as c:
            await c.execute(
                "INSERT INTO paper_positions(status,current_net_usd,closed_at) VALUES('CLOSED',-2,?)",
                (time.time(),),
            )
            await c.commit()
        ledger = Ledger(50)
        risk = RiskGuard(50)
        await restore(d.path, ledger, risk)
        assert ledger.equity == 48 and not risk.can_open_paper()

    asyncio.run(go())


def test_shared_capital_prevents_second_strategy_overallocation(tmp_path):
    async def go():
        d = Diary(str(tmp_path / "a.db"))
        await d.init()
        core = PaperEngine(d, capital=15)
        secondary = Engine(capital=15, reserved=lambda: core.used_capital)
        core.external_reserved = lambda: secondary.used_capital
        assert secondary.open(sf_op())
        op = dict(
            symbol="Y",
            buy="a",
            sell="b",
            notional=5,
            entry_buy=100,
            entry_sell=110,
            executable=10,
        )
        assert not await core.open(op)

    asyncio.run(go())


def test_dex_research_requires_key_and_valid_input_without_network():
    async def go():
        assert (await Provider("").price(1, "bad", "bad", 1))[
            "reason"
        ] == "DEX_API_KEY_MISSING"
        assert not (await Provider("key").price(1, "bad", "bad", 1))["ok"]

    asyncio.run(go())


def test_main_wires_secondary_runtime_and_cleans_up_without_network(
    tmp_path, monkeypatch
):
    async def go():
        import app.main as main
        from dataclasses import replace
        from app.live_trade_store import Store
        from app.strategy_runtime import StrategyRuntime

        events = []
        config = replace(
            main.config,
            token="offline-test",
            admin_id=1,
            db_path=str(tmp_path / "a.db"),
            runtime_state_path=str(tmp_path / "runtime.json"),
        )
        diary = Diary(config.db_path)

        class FakeScanner:
            clients = {}
            specs = {}
            funding = None
            universe = SimpleNamespace(symbols=[])

            async def start(self):
                events.append("scanner-start")

            async def close(self):
                events.append("scanner-close")

        class Runtime:
            services = {
                "spot_futures": SimpleNamespace(
                    service=SimpleNamespace(source=SimpleNamespace())
                ),
                "spot_spot": SimpleNamespace(source=SimpleNamespace()),
            }

            async def start(self):
                events.append("secondary-start")

        bundle = SimpleNamespace(
            runtime=Runtime(),
            sf_paper=Engine(),
            ss_paper=SimpleNamespace(used_capital=0),
            funding_paper=None,
        )

        async def close():
            events.append("secondary-close")

        bundle.close = close

        async def build(*args, **kwargs):
            events.append("secondary-build")
            return bundle

        async def boot(readers):
            return SimpleNamespace(ready=False, snapshot={})

        async def poll(bot):
            assert "secondary-start" in events
            assert main.live_stop.stopped and not main.live_supervisor.private_verified
            assert "БД LIVE" in main.startup_text

        class Session:
            async def close(self):
                events.append("bot-close")

        monkeypatch.setattr(main, "config", config)
        monkeypatch.setattr(main, "diary", diary)
        monkeypatch.setattr(main, "durable", Store(config.db_path))
        monkeypatch.setattr(main, "paper", PaperEngine(diary))
        monkeypatch.setattr(main, "bankroll", Ledger(50))
        monkeypatch.setattr(main, "scanner", FakeScanner())
        monkeypatch.setattr(main, "strategy_runtime", StrategyRuntime())
        monkeypatch.setattr(main, "build_bundle", build)
        monkeypatch.setattr(main, "bootstrap", boot)
        monkeypatch.setattr(main, "build_private_readers", lambda: ({}, {}))
        monkeypatch.setattr(
            main, "Bot", lambda **kw: SimpleNamespace(session=Session())
        )
        monkeypatch.setattr(main.dp, "start_polling", poll)
        await main.main()
        assert main.scanner.book_recorder.closed
        assert main.scanner.book_recorder.task is None
        assert events == [
            "scanner-start",
            "secondary-build",
            "secondary-start",
            "secondary-close",
            "scanner-close",
            "bot-close",
        ]

    asyncio.run(go())


def test_ccxt_requires_actual_usd_fees_and_uses_stable_short_client_id():
    from app.ccxt_executor import CCXTExecutor, client_id

    class Client:
        pass

    ex = CCXTExecutor("a", Client())
    for row in (
        {"filled": 1, "average": 100},
        {"filled": 1, "average": 100, "fee": {"cost": 0.001, "currency": "BNB"}},
    ):
        try:
            ex._result(row)
        except RuntimeError:
            pass
        else:
            raise AssertionError("unknown fee accepted as zero USD")
    value = "a" * 32 + ":entry-long"
    assert len(client_id(value)) == 32 and ":" not in client_id(value)
    assert client_id(value) == client_id(value) and client_id(value) != client_id(
        value + "x"
    )
